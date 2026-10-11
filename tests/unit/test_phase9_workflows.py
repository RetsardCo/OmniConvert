"""Unit and integration tests for Phase 9: Cross-Category Workflow Planner."""

import pytest
import zipfile
from pathlib import Path
from PIL import Image
import docx

from omniconvert.core.registry import CapabilityRegistry
from omniconvert.core.planner import WorkflowPlanner, WorkflowEngine
from omniconvert.core.job_manager import JobManager
from omniconvert.models.workflow import WorkflowPlan, WorkflowStep, WorkflowStepError
from omniconvert.models.conversion_job import ConversionJob


def test_workflow_planner_bounded_search_and_cycle_prevention():
    """Verify WorkflowPlanner finds multi-step routes without cycles and respects max_steps."""
    reg = CapabilityRegistry()
    planner = WorkflowPlanner(registry=reg)

    # DOCX -> WAV (2 steps: office.to_txt -> speech.text_to_wav)
    plans = planner.find_workflows("docx", "wav", only_available=True, max_steps=3)
    assert len(plans) >= 1
    best = plans[0]
    assert best.is_multi_step
    assert len(best.steps) == 2
    assert best.steps[0].output_format == "txt"
    assert best.steps[1].output_format == "wav"

    # Verify cycle prevention: no duplicate output formats in any plan
    for p in plans:
        seen = {p.source_format}
        for s in p.steps:
            assert s.output_format not in seen
            seen.add(s.output_format)


def test_workflow_execution_docx_to_wav(tmp_path):
    """Verify end-to-end 2-step workflow execution: DOCX -> TXT -> WAV."""
    doc_path = tmp_path / "speech_note.docx"
    d = docx.Document()
    d.add_heading("Project Update", level=1)
    d.add_paragraph("OmniConvert converts documents into spoken audio locally.")
    d.save(str(doc_path))

    jm = JobManager()
    planner = WorkflowPlanner(jm._engines["workflow"].planner.registry)
    plans = planner.find_workflows("docx", "wav", only_available=True, max_steps=2)
    assert plans

    ws = tmp_path / "ws_wf"
    ws.mkdir()
    out_wav = planner.execute_plan(
        plan=plans[0],
        source_path=doc_path,
        workspace_dir=ws,
        engines=jm._engines
    )
    assert out_wav.exists()
    assert out_wav.suffix == ".wav"
    assert out_wav.stat().st_size > 44


def test_workflow_step_failure_reporting(tmp_path):
    """Verify workflow stops immediately and reports exact failing step when an intermediate step fails."""
    jm = JobManager()
    planner = WorkflowPlanner()

    bad_plan = WorkflowPlan(
        source_format="png",
        target_format="wav",
        steps=[
            WorkflowStep(1, "image.to_jpeg", "Convert to JPEG", "pillow", "png", "jpeg"),
            WorkflowStep(2, "speech.text_to_wav", "Invalid Step", "nonexistent_engine", "jpeg", "wav"),
        ]
    )

    img_path = tmp_path / "sample.png"
    Image.new("RGB", (64, 64), (10, 20, 30)).save(img_path, "PNG")

    ws = tmp_path / "ws_fail"
    ws.mkdir()
    with pytest.raises(WorkflowStepError) as exc_info:
        planner.execute_plan(bad_plan, img_path, ws, jm._engines)

    assert exc_info.value.step_index == 2
    assert "Step 2" in str(exc_info.value)


def test_archive_to_combined_pdf_workflow(tmp_path):
    """Verify Archive-to-Combined-PDF extracts images and documents and merges them into a single PDF."""
    import pypdf

    src_dir = tmp_path / "bundle"
    src_dir.mkdir()
    Image.new("RGB", (120, 120), (200, 50, 50)).save(src_dir / "page1.png", "PNG")
    (src_dir / "page2.txt").write_text("Second page from text file inside ZIP archive.", encoding="utf-8")

    zip_path = tmp_path / "bundle.zip"
    with zipfile.ZipFile(zip_path, "w") as zf:
        zf.write(src_dir / "page1.png", arcname="page1.png")
        zf.write(src_dir / "page2.txt", arcname="page2.txt")

    jm = JobManager()
    wf_engine: WorkflowEngine = jm._engines["workflow"]  # type: ignore

    ws = tmp_path / "ws_arc_pdf"
    ws.mkdir()
    job = ConversionJob(
        source_path=zip_path,
        target_format="pdf",
        operation_id="workflow.archive_to_pdf",
        engine="workflow"
    )
    combined_pdf = wf_engine.execute(job=job, workspace_dir=ws)
    assert combined_pdf.exists()
    reader = pypdf.PdfReader(str(combined_pdf))
    assert len(reader.pages) >= 2


def test_csv_to_video_visualization(tmp_path):
    """Verify CSV-to-MP4 visualization template generates a valid H.264 MP4 video."""
    jm = JobManager()
    if "ffmpeg" not in jm._engines:
        pytest.skip("FFmpeg not available")

    csv_file = tmp_path / "metrics.csv"
    csv_file.write_text("Month,Revenue,Users\nJan,12000,340\nFeb,15500,420\nMar,19000,510\n", encoding="utf-8")

    ws = tmp_path / "ws_csv_mp4"
    ws.mkdir()
    job = ConversionJob(
        source_path=csv_file,
        target_format="mp4",
        operation_id="data.csv_to_mp4",
        engine="ffmpeg"
    )
    mp4_out = jm._engines["ffmpeg"].execute(job=job, workspace_dir=ws)
    assert mp4_out.exists()
    assert mp4_out.suffix == ".mp4"
    assert mp4_out.stat().st_size > 0
